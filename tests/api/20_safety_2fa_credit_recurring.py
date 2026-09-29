# Live exchange rates, the safety net (backup status, alerts), two-step sign-in (TOTP + recovery codes),
# credit notes, receipts, recurring invoices and the accountant's monthly pack.
import subprocess, time, json, os, urllib.request, urllib.error, sys, atexit, threading, hmac, hashlib, base64, struct, http.client, io, datetime
from http.server import BaseHTTPRequestHandler, HTTPServer
DB = os.environ.get('TEST_DATABASE_URL', 'postgres://test@127.0.0.1:5439/gotham_test')
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
os.chdir(os.path.join(ROOT, 'apps', 'crm'))
class Mock(BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def do_GET(self):
        b = json.dumps({'rates': {'GHS': 15.5, 'EUR': 0.9, 'GBP': 0.78}, 'time_last_update_utc': 'Tue, 29 Sep 2026 00:02:31 +0000'}).encode()
        self.send_response(200); self.send_header('content-type', 'application/json'); self.end_headers(); self.wfile.write(b)
api = HTTPServer(('127.0.0.1', 4294), Mock); threading.Thread(target=api.serve_forever, daemon=True).start()
env = dict(os.environ); env.update({'DATABASE_URL': DB, 'SESSION_SECRET': 't', 'PORT': '4199', 'ADMIN_EMAIL': 'owner@gotham.test', 'ADMIN_PASSWORD': 'ownerpass123', 'ADMIN_NAME': 'Owner',
    'MEDIA_DIR': '/tmp/mediatest', 'PUBLIC_URL': 'http://127.0.0.1:4199', 'SITE_URL': 'http://example.org', 'FX_API': 'http://127.0.0.1:4294/fx', 'JOBS_EVERY_MS': '2500'})
os.makedirs('/tmp/mediatest', exist_ok=True)
srv = subprocess.Popen(['node', 'server.js'], env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
atexit.register(lambda: srv.poll() is None and srv.terminate())
BASE = 'http://127.0.0.1:4199'
def call(path, data=None, method=None, cookie=None, headers=None, text=False, binary=False):
    req = urllib.request.Request(BASE + path, data=json.dumps(data).encode() if data is not None else None, method=method or ('POST' if data is not None else 'GET'))
    req.add_header('content-type', 'application/json')
    for k, v in (headers or {}).items(): req.add_header(k, v)
    if cookie: req.add_header('cookie', cookie)
    try:
        r = urllib.request.urlopen(req, timeout=60); b = r.read()
        return r.status, (b if binary else b.decode() if text else json.loads(b or b'{}')), r.headers
    except urllib.error.HTTPError as e:
        b = e.read()
        try: return e.code, (b.decode() if text else json.loads(b or b'{}')), e.headers
        except Exception: return e.code, b[:300], e.headers
for _ in range(60):
    try:
        if call('/api/health')[0] == 200: break
    except Exception: pass
    time.sleep(0.4)
else: print('NO START'); print(srv.stdout.read()[:3000]); sys.exit(1)
def psql(sql): return subprocess.run(['psql', DB, '-qAtc', sql], capture_output=True, text=True, check=True).stdout.strip()
ok = lambda m: print('  ok', m)
psql('update users set must_change=false')
IP = iter(range(10, 250))
def login(email, pw, code=None):
    st, d, h = call('/api/login', {'email': email, 'password': pw}, headers={'x-real-ip': '10.66.%d.1' % next(IP)})
    return st, d, (h.get('set-cookie') or '').split(';')[0]
st, _, ck = login('owner@gotham.test', 'ownerpass123'); assert st == 200 and ck, st

# ---------- live exchange rates ----------
st, R, h = call('/api/rates'); assert st == 200 and R['rates'] == {'USD': 15.5, 'EUR': 17.2222, 'GBP': 19.8718} and R['updated'].startswith('2026-09-29') and 'Bank of Ghana' in R['source'], R
assert 'max-age' in (h.get('cache-control') or '')
assert json.loads(psql("select value from settings where key='fx_cache'"))['rates']['USD'] == 15.5
ok('the site gets live cedi rates for USD, EUR and GBP (cached hourly, kept through restarts, labelled as market rate vs Bank of Ghana)')

# ---------- safety net ----------
st, S, _ = call('/api/safety', cookie=ck); assert st == 200 and any(k == 'backup' for k, _ in S['problems']), S
st, e, _ = call('/api/safety', {'alert_emails': 'not-an-email'}, cookie=ck); assert st == 400 and e['error'] == 'bad_email', e
st, S, _ = call('/api/safety', {'alert_emails': 'ops@gotham.test, owner@gotham.test', 'alert_whatsapp': '+233 55 000 0000'}, cookie=ck); assert S['alert_emails'] == 'ops@gotham.test, owner@gotham.test' and S['alert_whatsapp'] == '+233550000000', S
st, T, _ = call('/api/safety', {'test': True}, cookie=ck); assert T['sent_to'] == 2, T
assert psql("select count(*) from outbox where subject='Gotham platform: test alert' and kind='alert'") == '2'
now = datetime.datetime.utcnow().strftime('%Y-%m-%dT%H:%M:%SZ')
psql("insert into settings(key,value) values ('backup_status','{\"ok\":true,\"at\":\"%s\",\"file\":\"x.tar\",\"bytes\":12345,\"detail\":\"restore test passed: 1 users\"}') on conflict (key) do update set value=excluded.value" % now)
st, S, _ = call('/api/safety', cookie=ck); assert S['problems'] == [] and S['backup']['ok'] is True, S
psql("update settings set value='{\"ok\":false,\"at\":\"%s\",\"detail\":\"database dump failed\"}' where key='backup_status'" % now)
psql("delete from settings where key in ('watchdog_at','alert_last_backup')")
for _ in range(20):
    if psql("select count(*) from outbox where subject='Gotham platform: please check' and body like '%%database dump failed%%'") == '2': break
    time.sleep(0.5)
else: raise AssertionError('watchdog did not warn')
time.sleep(3)
assert psql("select count(*) from outbox where subject='Gotham platform: please check' and body like '%%database dump failed%%'") == '2', 'no repeat within 12 hours'
st, x, _ = call('/api/safety'); assert st == 401
ok('safety net: backup status and failures are watched hourly; warnings go to the chosen people once per 12 hours; test warning works')

# ---------- two-step sign-in ----------
def b32d(s): return base64.b32decode(s + '=' * (-len(s) % 8))
def totp(secret, t=None):
    c = int((t or time.time()) // 30); h = hmac.new(b32d(secret), struct.pack('>Q', c), hashlib.sha1).digest(); o = h[-1] & 15
    return '%06d' % ((struct.unpack('>I', h[o:o + 4])[0] & 0x7fffffff) % 1000000)
st, d, _ = call('/api/users', {'name': 'Ama', 'email': 'ama@gotham.test', 'role': 'admin'}, cookie=ck); pw = d['password']
psql("update users set must_change=false where email='ama@gotham.test'")
st, _, ack = login('ama@gotham.test', pw); assert st == 200
st, g, _ = call('/api/2fa', cookie=ack); assert g == {'on': False, 'recovery_left': 0}, g
st, su, _ = call('/api/2fa/setup', {}, cookie=ack); assert st == 200 and su['uri'].startswith('otpauth://totp/') and '<svg' in su['qr'], su
assert psql("select totp_secret from users where email='ama@gotham.test'") != su['secret'], 'secret stored encrypted'
st, e, _ = call('/api/2fa/enable', {'code': '000000'}, cookie=ack); assert st == 400 and e['error'] == 'bad_code'
st, en, _ = call('/api/2fa/enable', {'code': totp(su['secret'])}, cookie=ack); assert st == 200 and len(en['recovery_codes']) == 8, en
st, d, c0 = login('ama@gotham.test', pw); assert st == 200 and d['need_code'] is True and d['ticket'] and not c0.startswith('gc_session=') or c0 == '', (d, c0)
t1 = d['ticket']
st, e, _ = call('/api/login/code', {'ticket': t1, 'code': '123456'}, headers={'x-real-ip': '10.67.1.1'}); assert st == 401 and e['error'] == 'bad_code'
time.sleep(max(0, 30 - time.time() % 30) + 0.5)   # a fresh time step: the enable code cannot be replayed
code = totp(su['secret'])
st, me, h = call('/api/login/code', {'ticket': t1, 'code': code}, headers={'x-real-ip': '10.67.1.2'}); assert st == 200 and me['email'] == 'ama@gotham.test' and 'gc_session=' in h.get('set-cookie', ''), me
st, d, _ = login('ama@gotham.test', pw)
st, e, _ = call('/api/login/code', {'ticket': d['ticket'], 'code': code}, headers={'x-real-ip': '10.67.1.3'}); assert st == 401, 'the same code cannot be used twice'
rc = en['recovery_codes'][0]
st, me, _ = call('/api/login/code', {'ticket': d['ticket'], 'code': rc.lower()}, headers={'x-real-ip': '10.67.1.4'}); assert st == 200, me
st, d, _ = login('ama@gotham.test', pw)
st, e, _ = call('/api/login/code', {'ticket': d['ticket'], 'code': rc}, headers={'x-real-ip': '10.67.1.5'}); assert st == 401, 'a recovery code works once'
st, g, _ = call('/api/2fa', cookie=ack); assert g == {'on': True, 'recovery_left': 7}, g
st, e, _ = call('/api/login/code', {'ticket': 'forged.ticket', 'code': '123456'}, headers={'x-real-ip': '10.67.1.6'}); assert st == 401 and e['error'] == 'expired'
st, U, _ = call('/api/users', cookie=ck); assert [u for u in U['rows'] if u['email'] == 'ama@gotham.test'][0]['totp_on'] is True
st, _, _ = call('/api/users/ama%40gotham.test', {'reset_2fa': True}, method='PATCH', cookie=ck); assert st == 200
st, d, c2 = login('ama@gotham.test', pw); assert st == 200 and 'need_code' not in d, d
st, x, _ = call('/api/me', cookie=ack); assert st == 401, 'resetting two-step signs the person out'
ok('two-step sign-in: QR setup, encrypted secret, code after password, no replay, one-time recovery codes, owner can reset it')

# ---------- credit notes and receipts ----------
call('/api/invoice-settings', {'inv_vat_registered': 'yes', 'inv_tin': 'C0012345678'}, cookie=ck)
def issue(lines, email='buyer@b.test', name='Buyer Ltd'):
    st, d, _ = call('/api/invoices', {'bill_to': {'name': name, 'email': email}, 'lines': lines}, cookie=ck); assert st == 200, d
    st, x, _ = call('/api/invoices/%d/issue' % d['id'], {}, cookie=ck); assert st == 200, x
    return d['id'], x['link'].rsplit('/', 1)[1]
i1, t1 = issue([{'description': 'Advisory', 'qty': 1, 'unit': 1000}])
st, e, _ = call('/api/invoices/%d/credit' % i1, {}, cookie=ck); assert e['error'] == 'reason_required'
st, e, _ = call('/api/invoices/%d/credit' % i1, {'reason': 'x', 'lines': [{'description': 'too much', 'qty': 1, 'unit': 5000}]}, cookie=ck); assert st == 400 and e['error'] == 'credit_too_big' and e['left'] == 1200, e
st, c, _ = call('/api/invoices/%d/credit' % i1, {'reason': 'Fee reduced as agreed', 'lines': [{'description': 'Discount agreed', 'qty': 1, 'unit': 250}]}, cookie=ck)
Y = time.strftime('%Y'); assert st == 200 and c['number'] == 'CRN-%s-0001' % Y and c['status'] == 'sent', c
cn = json.loads(psql("select row_to_json(i) from invoices i where id=%d" % c['id'])); assert float(cn['total']) == 300 and float(cn['vat']) == 37.5 and cn['credit_for'] == i1, cn
st, one, _ = call('/api/invoices/%d' % i1, cookie=ck); assert one['credits'][0]['number'] == c['number']
st, L, _ = call('/api/invoices', cookie=ck); r1 = [r for r in L['rows'] if r['id'] == i1][0]; assert r1['balance'] == 900 and float(r1['credited']) == 300, r1
st, html, _ = call('/i/' + t1, text=True); assert 'Credit note CRN-%s-0001' % Y in html and 'GHS 900.00' in html
st, html, _ = call('/i/' + cn['token'], text=True); assert 'CREDIT NOTE' in html and 'Credit against invoice' in html and 'Credited to' in html and 'How to pay' not in html
# pay the rest: receipt
st, p, _ = call('/api/invoices/%d/payment' % i1, {'amount': 900, 'method': 'momo', 'reference': 'MP1'}, cookie=ck); assert p['status'] == 'paid', p
rno = psql("select receipt_no from invoice_payments where invoice_id=%d" % i1); assert rno == 'RCT-%s-0001' % Y, rno
assert psql("select count(*) from outbox where to_email='buyer@b.test' and subject like 'Receipt RCT-%%'") == '1'
pid = psql("select id from invoice_payments where invoice_id=%d" % i1)
st, rh, _ = call('/i/%s/receipt/%s' % (t1, pid), text=True); assert st == 200 and 'RECEIPT' in rh and rno in rh and 'GHS 900.00' in rh and 'GHS 0.00' in rh
st, x, _ = call('/i/%s/receipt/999999' % t1, text=True); assert st == 404
st, x, _ = call('/api/invoices/%d/receipt/%s' % (i1, pid), {'to': 'accounts@b.test'}, cookie=ck); assert st == 200
assert psql("select count(*) from outbox where to_email='accounts@b.test'") == '1'
# whole invoice credited
i2, t2 = issue([{'description': 'Workshop', 'qty': 2, 'unit': 500}])
st, c2, _ = call('/api/invoices/%d/credit' % i2, {'reason': 'Duplicate invoice'}, cookie=ck); assert c2['status'] == 'credited', c2
st, html, _ = call('/i/' + t2, text=True); assert 'CREDITED' in html and 'How to pay' not in html
st, csv, _ = call('/api/invoices/vat.csv?month=' + time.strftime('%Y-%m'), cookie=ck, text=True)
rows = csv.lstrip('﻿').splitlines(); crn = [r for r in rows if r.startswith('CRN-')]; assert len(crn) == 2 and ',credit note,' in crn[0] and ',-250,' in crn[0] and ',-37.5,' in crn[0], crn
ok('credit notes CRN-YYYY-NNNN: reduce or cancel an issued invoice (never more than is left), tax goes down with them, shown on the client page; receipts RCT-YYYY-NNNN with page and email')

# ---------- recurring invoices ----------
i3, _ = issue([{'description': 'Investor desk retainer, {month}', 'qty': 1, 'unit': 11500}], email='desk@c.test', name='Nordic Wind AS')
today = time.strftime('%Y-%m-%d')
st, e, _ = call('/api/invoice-schedules', {'from_invoice_id': i3, 'every': 'month', 'next_date': '2020-01-01'}, cookie=ck); assert e['error'] == 'date_in_past'
st, sc, _ = call('/api/invoice-schedules', {'from_invoice_id': i3, 'every': 'month', 'next_date': today, 'auto_send': True}, cookie=ck); assert st == 200, sc
for _ in range(30):
    if psql("select runs from invoice_schedules where id=%d" % sc['id']) == '1': break
    time.sleep(0.5)
else: raise AssertionError('schedule did not run')
new = json.loads(psql("select row_to_json(i) from invoices i where created_by='schedule' order by id desc limit 1"))
month = datetime.date.today().strftime('%B %Y')
assert new['number'] and new['status'] == 'sent' and new['lines'][0]['description'] == 'Investor desk retainer, ' + month and float(new['total']) == 13800, new
assert psql("select count(*) from outbox where to_email='desk@c.test' and subject like 'Invoice %s from%%'" % new['number']) == '1'
nd = psql("select next_date from invoice_schedules where id=%d" % sc['id']); assert nd > today, nd
st, S2, _ = call('/api/invoice-schedules', cookie=ck); assert S2['rows'][0]['runs'] == 1 and S2['rows'][0]['last_number'] == new['number'] and S2['rows'][0]['amount'] == 13800, S2
st, _, _ = call('/api/invoice-schedules/%d' % sc['id'], {'active': False}, cookie=ck); assert psql("select active from invoice_schedules where id=%d" % sc['id']) == 'f'
# foreign currency: a draft and a task instead
st, d, _ = call('/api/invoices', {'bill_to': {'name': 'Euro Client', 'residency': 'non_resident'}, 'currency': 'EUR', 'fx_rate': 17, 'tax_mode': 'zero', 'tax_note': 'Export', 'lines': [{'description': 'Retainer {month}', 'qty': 1, 'unit': 800}]}, cookie=ck)
call('/api/invoices/%d/issue' % d['id'], {}, cookie=ck)
st, sc2, _ = call('/api/invoice-schedules', {'from_invoice_id': d['id'], 'every': 'quarter', 'next_date': today}, cookie=ck)
for _ in range(30):
    if psql("select runs from invoice_schedules where id=%d" % sc2['id']) == '1': break
    time.sleep(0.5)
assert psql("select count(*) from invoices where created_by='schedule' and currency='EUR' and number is null") == '1'
assert psql("select count(*) from tasks where title like 'Recurring invoice for Euro Client is ready as a draft%%'") == '1'
ok('recurring invoices: monthly/quarterly/yearly copies with {month} filled in; cedi ones issued and emailed, foreign-currency ones wait as drafts for the day\'s rate; pause/resume')

# ---------- accountant pack ----------
st, xl, h = call('/api/invoices/pack.xlsx?month=' + time.strftime('%Y-%m'), cookie=ck, binary=True); assert st == 200 and 'spreadsheetml' in h['content-type']
import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(xl))
assert wb.sheetnames == ['Summary', 'Invoices & credit notes', 'Payments & WHT', 'Owed at month end', 'Xero import', 'QuickBooks import'], wb.sheetnames
summ = {r[0].value: r[1].value for r in wb['Summary'].iter_rows(min_row=2) if r[0].value and len(r) > 1}
assert summ['Credit notes issued'] == 2 and summ['Invoices issued'] >= 4, summ
assert wb['Payments & WHT'].max_row >= 2 and wb['Xero import'].cell(1, 1).value == '*ContactName'
st, xc, _ = call('/api/invoices/xero.csv?month=' + time.strftime('%Y-%m'), cookie=ck, text=True)
assert xc.lstrip('﻿').startswith('*ContactName,EmailAddress') and ',-1,' in xc, xc[:300]
st, qc, _ = call('/api/invoices/quickbooks.csv?month=' + time.strftime('%Y-%m'), cookie=ck, text=True); assert qc.lstrip('﻿').startswith('InvoiceNo,Customer'), qc[:100]
st, x, _ = call('/api/invoices/pack.xlsx'); assert st == 401
ok('accountant pack: Excel with summary, invoices and credit notes, payments/WHT, owed at month end, Xero and QuickBooks import sheets; CSVs for both')
print('PASSED')
